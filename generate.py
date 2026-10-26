# -*- coding: utf-8 -*-
"""
खाजगी विहीर अधिग्रहण प्रस्ताव — Excel मधून प्रत्येक व्यक्तीसाठी स्वतंत्र PDF
Private Well Acquisition Proposal — one PDF per Excel row

The page layout lives in web/form.js (shared with the browser app,
Vihir_Adhigrahan_App.html); this script reads Excel/CSV rows and prints
each one to a vector PDF with Chromium.

Usage:
    python generate.py                      -> reads Vihir_Data.xlsx, writes PDFs to output/
    python generate.py MyList.xlsx          -> reads a different Excel file
    python generate.py Vihir_Proposals.csv  -> reads the CSV exported from the app
    python generate.py --blank              -> writes the clean blank form (Blank_Form.pdf)

Needs (one time):
    pip install openpyxl playwright
    python -m playwright install chromium
"""
import sys, re, csv, base64, datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
FONT_DIR = BASE / "fonts"
WEB_DIR = BASE / "web"
OUT_DIR = BASE / "output"

# ---------------------------------------------------------------- columns
# (Excel header, key, example)  — the header text is what the script reads.
COLUMNS = [
    ("अ.क्र.", "sr", "1"),
    ("जिल्हा", "district", "जालना"),
    ("तालुका", "taluka", "जालना"),
    ("जि.प. उप विभाग", "upvibhag", ""),
    ("गावाचे नाव (मौजे)", "gav", "निधोना"),
    ("ग्रामपंचायत", "gp", "निधोना"),
    ("गाव/वाडी/तांडा/वस्ती (जेथे पाणी पुरवठा)", "wadi", "आंबेडकर नगर"),
    ("तलाठी सज्जा", "sajja", "निधोना"),
    ("विहीर मालकाचे नाव", "owner", "रामराव विठ्ठल पाटील"),
    ("मालक राहणार", "owner_ra", "निधोना"),
    ("गट क्र.", "gat", "145"),
    ("प्रस्ताव दिनांक", "dinank", "15/01/2027"),
    ("ग्रामसभा ठराव क्र.", "tharav_no", "4"),
    ("ठराव दिनांक", "tharav_dinank", "10/01/2027"),
    ("ग्रामसभा सन", "sabha_san", "2026-27"),
    ("पाहणी दिनांक", "pahani_dinank", "12/01/2027"),
    ("करारनामा दिनांक", "karar_dinank", "14/01/2027"),
    ("टंचाई माहे ... पासून", "tanchai_mahe", "जानेवारी 2027"),
    ("लोकसंख्या सन 2011", "pop_2011", "464"),
    ("आजची लोकसंख्या", "pop_today", "520"),
    ("विहिरीचे गावापर्यंत अंतर (कि.मी.)", "antar", "1.2"),
    ("उपाय योजना", "upay", "खाजगी विहीर अधिग्रहण"),
    ("पंच 1", "panch1", "सुभाष रामा जाधव"),
    ("पंच 2", "panch2", "शेख अहमद शेख रशीद"),
    ("पंच 3", "panch3", "गणेश बाबुराव काळे"),
    ("पंच 4", "panch4", ""),
    ("पंच 5", "panch5", ""),
    ("ठराव सूचक", "suchak", "सुभाष रामा जाधव"),
    ("ठराव अनुमोदक", "anumodak", "गणेश बाबुराव काळे"),
    ("संयुक्त पाहणी अधिकारी नाव", "officer_name", ""),
    ("संयुक्त पाहणी अधिकारी पदनाम", "officer_pad", ""),
    ("क्षेत्रीय अधिकारी (पा.टंचाई) नाव", "kshetriya", ""),
    ("सरपंच नाव", "sarpanch", ""),
    ("ग्रामसेवक/ग्रा.वि.अ. नाव", "gramsevak", ""),
    ("पोलीस पाटील", "police_patil", ""),
    ("सहकारी संस्था चेअरमन", "chairman", ""),
    ("ग्रा.पं. महिला सदस्य", "mahila", ""),
    ("अंगणवाडी सेविका", "anganwadi", ""),
    ("शाळेचे मुख्याध्यापक", "mukhyadhyapak", ""),
    ("मंडळ अधिकारी / तलाठी नाव", "talathi", ""),
    ("सार्व. विहीर/बोअर संख्या", "sv_count", "1"),
    ("सार्व. विहीर सद्यस्थिती", "sv_status", "कोरडी"),
    ("एकूण हातपंप", "hp_total", "2"),
    ("चालू हातपंप", "hp_chalu", "0"),
    ("थांबून चालणारे हातपंप", "hp_thambun", "1"),
    ("पाण्याअभावी बंद हातपंप", "hp_band", "1"),
    ("कायम बंद हातपंप", "hp_kayam", "0"),
    ("विद्युत पंप संख्या", "vp_count", "1"),
    ("विद्युत पंप सद्यस्थिती", "vp_status", "पाण्याअभावी बंद"),
    ("नळ योजना संख्या", "nal_count", "1"),
    ("नळ योजना सद्यस्थिती", "nal_status", "उद्भव कोरडा"),
    ("मोबदला कोणाच्या नावे", "mobadla", ""),
    ("सहहिस्सेदार 1", "sah1", ""),
    ("सहहिस्सेदार 2", "sah2", ""),
    ("सहहिस्सेदार 3", "sah3", ""),
    ("सहहिस्सेदार 4", "sah4", ""),
    ("सहहिस्सेदार 5", "sah5", ""),
    ("जलयुक्त शिवार कामे", "jalyukt", "निरंक"),
    ("टंचाई आराखडा कालावधी", "ar_kal", "जानेवारी 2027 ते जून 2027"),
    ("आराखड्यात समाविष्ट (आहे/नाही)", "ar_samavisht", "आहे"),
    ("जाने-मार्च प्रस्तावित उपाय", "ar_jm_upay", "निरंक"),
    ("जाने-मार्च खर्च (लाखात)", "ar_jm_kharch", "0.00"),
    ("एप्रिल-जून प्रस्तावित उपाय", "ar_aj_upay", "विहिर अधिग्रहण करणे"),
    ("एप्रिल-जून खर्च (लाखात)", "ar_aj_kharch", "0.54"),
]
HEADER_TO_KEY = {h: k for h, k, _ in COLUMNS}
KEYS = {k for _, k, _ in COLUMNS}
REQUIRED = ["gav", "owner", "owner_ra", "gat", "dinank"]
LABEL = {k: h for h, k, _ in COLUMNS}
# Older Excel files have no जिल्हा/तालुका columns; they were made for ता. जालना.
DEFAULTS = {"district": "जालना", "taluka": "जालना"}


# ---------------------------------------------------------------- values
def fmt(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:g}"
    return str(v).strip()


def problems(d):
    """Missing required values and numbers that are not numbers."""
    out = [LABEL[k] for k in REQUIRED if not d.get(k)]
    for k in ("ar_jm_kharch", "ar_aj_kharch"):
        if d.get(k):
            try:
                float(d[k])
            except ValueError:
                out.append(f"{LABEL[k]} = '{d[k]}' (आकडा नाही)")
    return out


# ---------------------------------------------------------------- input
def rows_to_dicts(header, body):
    keys = [HEADER_TO_KEY.get(h) or (h if h in KEYS else None) for h in header]
    unknown = [h for h, k in zip(header, keys) if h and not k]
    if unknown:
        print("⚠ ओळखता न आलेले कॉलम (दुर्लक्षित):", ", ".join(unknown))
    out = []
    for r in body:
        d = {k: fmt(v) for k, v in zip(keys, r) if k}
        if d.get("gav") or d.get("owner"):
            for k, v in DEFAULTS.items():
                d[k] = d.get(k) or v
            out.append(d)
    return out


def read_excel(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    return rows_to_dicts([fmt(h) for h in rows[0]], rows[1:])


def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return []
    return rows_to_dicts([h.strip() for h in rows[0]], rows[1:])


def safe(s):
    s = re.sub(r'[\\/:*?"<>|]+', "", s or "").strip()
    return re.sub(r"\s+", "_", s)[:60]


# ---------------------------------------------------------------- render
def font_face():
    out = []
    for w, f in ((400, "Mukta-Regular.ttf"), (600, "Mukta-SemiBold.ttf"), (700, "Mukta-Bold.ttf")):
        b64 = base64.b64encode((FONT_DIR / f).read_bytes()).decode()
        out.append(f"@font-face{{font-family:'Mukta';font-weight:{w};src:url(data:font/ttf;base64,{b64}) format('truetype')}}")
    return "\n".join(out)


def html_to_pdf(jobs):
    """jobs: [(data_dict, pdf_path)] — each rendered by web/form.js inside Chromium."""
    from playwright.sync_api import sync_playwright
    form_js = (WEB_DIR / "form.js").read_text(encoding="utf-8")
    opts = {"fontCss": font_face(), "css": (WEB_DIR / "form.css").read_text(encoding="utf-8")}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        for data, pdf_path in jobs:
            page.set_content("<!doctype html><meta charset='utf-8'>")
            page.add_script_tag(content=form_js)
            doc = page.evaluate("([d, o]) => VihirForm.renderDocument(d, o)", [data, opts])
            page.set_content(doc, wait_until="load")
            page.evaluate("document.fonts.ready")
            page.pdf(path=str(pdf_path), prefer_css_page_size=True, print_background=True,
                     margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
            print("✔", Path(pdf_path).name)
        browser.close()


def main():
    args = sys.argv[1:]
    if "--blank" in args:
        html_to_pdf([({}, BASE / "Blank_Form.pdf")])
        return
    src = Path(args[0]) if args else BASE / "Vihir_Data.xlsx"
    rows = read_csv(src) if src.suffix.lower() == ".csv" else read_excel(src)
    if not rows:
        print("माहिती सापडली नाही:", src)
        return
    OUT_DIR.mkdir(exist_ok=True)
    jobs, used, incomplete = [], set(), 0
    for n, d in enumerate(rows, 1):
        sr = d.get("sr") or str(n)
        bad = problems(d)
        if bad:
            incomplete += 1
            print(f"⚠ ओळ {n} ({d.get('owner') or '—'}) अपूर्ण: " + ", ".join(bad))
        stem = f"{sr.zfill(3)}_{safe(d.get('gav'))}_{safe(d.get('owner'))}"
        name, k = stem, 2
        while name in used:  # never overwrite another row's PDF
            name, k = f"{stem}_{k}", k + 1
        used.add(name)
        jobs.append((d, OUT_DIR / f"{name}.pdf"))
    html_to_pdf(jobs)
    print(f"\nएकूण {len(jobs)} PDF तयार → {OUT_DIR}" + (f"  (⚠ {incomplete} अपूर्ण)" if incomplete else ""))


if __name__ == "__main__":
    main()
